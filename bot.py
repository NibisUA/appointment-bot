#!/usr/bin/env python3
# ================================================================
#  bot.py — Telegram-бот для запису на манікюр/педікюр
# ================================================================

import asyncio
import logging
import sqlite3
from datetime import datetime, timedelta, date
from calendar import monthrange

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    Message, CallbackQuery,
    InlineKeyboardMarkup, InlineKeyboardButton,
    FSInputFile
)
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from apscheduler.triggers.date import DateTrigger

from config import (
    BOT_TOKEN, ADMIN_ID, CHANNEL_ID, CHANNEL_LINK,
    SCHEDULE_CHANNEL_ID, DATABASE_NAME
)

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger(__name__)

EXCEL_FILE = "bookings.xlsx"

SERVICES = [
    ("man_french",  "💅 Манікюр — Французький", 1000),
    ("man_square",  "💅 Манікюр — Квадрат",       500),
    ("ped_classic", "🦶 Педікюр — Класичний",      800),
    ("ped_device",  "🦶 Педікюр — Апаратний",      900),
]
SERVICE_MAP = {s[0]: s for s in SERVICES}

REMINDER_OPTIONS = [
    ("30m",  "За 30 хвилин"),
    ("1h",   "За 1 годину"),
    ("2h",   "За 2 години"),
    ("3h",   "За 3 години"),
    ("6h",   "За 6 годин"),
    ("12h",  "За 12 годин"),
    ("24h",  "За 24 години"),
    ("off",  "🔕 Вимкнути"),
]
REMINDER_MINUTES = {
    "30m": 30, "1h": 60, "2h": 120,
    "3h": 180, "6h": 360, "12h": 720, "24h": 1440, "off": 0
}

# ================================================================
# FSM
# ================================================================
class Booking(StatesGroup):
    choose_service   = State()
    choose_date_a    = State()
    choose_time_a    = State()
    choose_date_b    = State()
    choose_time_b    = State()
    choose_service_b = State()
    enter_name       = State()
    enter_phone      = State()
    enter_comment    = State()
    confirm          = State()
    confirm_another  = State()   # підтвердження повторного запису

class Admin(StatesGroup):
    add_pick_date       = State()
    add_pick_times      = State()
    add_pick_services   = State()   # вибір послуг для кожного часу
    rm_pick_date        = State()
    rm_pick_time        = State()
    close_pick_date     = State()
    cancel_pick_date    = State()
    cancel_pick_booking = State()
    view_pick_date      = State()

# ================================================================
# БАЗА ДАНИХ
# ================================================================
def db():
    conn = sqlite3.connect(DATABASE_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with db() as c:
        c.executescript("""
            CREATE TABLE IF NOT EXISTS slots (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                slot_date TEXT NOT NULL,
                slot_time TEXT NOT NULL,
                is_booked        INTEGER DEFAULT 0,
                is_closed        INTEGER DEFAULT 0,
                allowed_services TEXT    DEFAULT '*',
                UNIQUE(slot_date, slot_time)
            );
            CREATE TABLE IF NOT EXISTS bookings (
                id                   INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id              INTEGER NOT NULL,
                username             TEXT,
                client_name          TEXT NOT NULL,
                phone                TEXT NOT NULL,
                slot_date            TEXT NOT NULL,
                slot_time            TEXT NOT NULL,
                service_id           TEXT DEFAULT 'man_french',
                comment              TEXT DEFAULT '',
                created_at           TEXT DEFAULT (datetime('now','localtime')),
                reminder_before_sent INTEGER DEFAULT 0,
                reminder_after_sent  INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS user_settings (
                user_id      INTEGER PRIMARY KEY,
                reminder_key TEXT DEFAULT '1h'
            );
        """)

    # Міграція для старих БД
    with db() as c:
        b_cols = [r[1] for r in c.execute("PRAGMA table_info(bookings)").fetchall()]
        for col, dflt in [
            ("reminder_before_sent", "0"),
            ("reminder_after_sent",  "0"),
            ("service_id",           "'man_french'"),
            ("comment",              "''"),
        ]:
            if col not in b_cols:
                c.execute(f"ALTER TABLE bookings ADD COLUMN {col} TEXT DEFAULT {dflt}")
                logger.info(f"Міграція bookings: {col}")

        s_cols = [r[1] for r in c.execute("PRAGMA table_info(slots)").fetchall()]
        if "allowed_services" not in s_cols:
            c.execute("ALTER TABLE slots ADD COLUMN allowed_services TEXT DEFAULT '*'")
            logger.info("Міграція slots: allowed_services")

        us_cols = [r[1] for r in c.execute("PRAGMA table_info(user_settings)").fetchall()]
        if "reminder_key" not in us_cols and us_cols:
            c.execute("ALTER TABLE user_settings ADD COLUMN reminder_key TEXT DEFAULT '1h'")
            logger.info("Міграція user_settings: reminder_key")

    logger.info("БД ініціалізована")

# ── Слоти ─────────────────────────────────────────────────────
def get_available_dates(service_id=None):
    """Дати з вільними слотами. Якщо service_id — тільки де є слоти для цієї послуги."""
    today = date.today().strftime("%Y-%m-%d")
    end   = (date.today() + timedelta(days=31)).strftime("%Y-%m-%d")
    with db() as c:
        rows = c.execute(
            "SELECT DISTINCT slot_date, allowed_services FROM slots "
            "WHERE slot_date>=? AND slot_date<=? AND is_booked=0 AND is_closed=0 "
            "ORDER BY slot_date", (today, end)
        ).fetchall()
    if not service_id:
        return list(dict.fromkeys(r["slot_date"] for r in rows))
    result = []
    seen   = set()
    for r in rows:
        sd  = r["slot_date"]
        svc = r["allowed_services"] or "*"
        if sd not in seen:
            if svc == "*" or service_id in [s.strip() for s in svc.split(",")]:
                result.append(sd)
                seen.add(sd)
    return result

def get_all_slot_dates():
    today = date.today().strftime("%Y-%m-%d")
    end   = (date.today() + timedelta(days=60)).strftime("%Y-%m-%d")
    with db() as c:
        rows = c.execute(
            "SELECT DISTINCT slot_date FROM slots "
            "WHERE slot_date>=? AND slot_date<=? ORDER BY slot_date", (today, end)
        ).fetchall()
    return [r["slot_date"] for r in rows]

def get_booking_dates():
    today = date.today().strftime("%Y-%m-%d")
    with db() as c:
        rows = c.execute(
            "SELECT DISTINCT slot_date FROM bookings "
            "WHERE slot_date>=? ORDER BY slot_date", (today,)
        ).fetchall()
    return [r["slot_date"] for r in rows]

def get_free_slots(slot_date, service_id=None):
    """Вільні слоти. Якщо service_id вказано — фільтрує по дозволених послугах."""
    with db() as c:
        rows = c.execute(
            "SELECT slot_time, allowed_services FROM slots "
            "WHERE slot_date=? AND is_booked=0 AND is_closed=0 "
            "ORDER BY slot_time", (slot_date,)
        ).fetchall()
    result = []
    for r in rows:
        if service_id:
            svc = r["allowed_services"] or "*"
            if svc != "*" and service_id not in [s.strip() for s in svc.split(",")]:
                continue
        result.append(r["slot_time"])
    return result

def get_all_slots(slot_date):
    with db() as c:
        return c.execute(
            "SELECT slot_time, is_booked, is_closed FROM slots "
            "WHERE slot_date=? ORDER BY slot_time", (slot_date,)
        ).fetchall()

def get_bookings_for_date(slot_date):
    with db() as c:
        return c.execute(
            "SELECT id, client_name, phone, slot_time, service_id, comment, username, user_id FROM bookings "
            "WHERE slot_date=? ORDER BY slot_time", (slot_date,)
        ).fetchall()

def get_user_active_bookings(uid):
    """Всі майбутні записи користувача."""
    today = date.today().strftime("%Y-%m-%d")
    with db() as c:
        return c.execute(
            "SELECT * FROM bookings WHERE user_id=? AND slot_date>=? ORDER BY slot_date, slot_time",
            (uid, today)
        ).fetchall()

def create_booking(uid, uname, name, phone, sd, st, service_id, comment=""):
    with db() as c:
        cur = c.execute(
            "INSERT INTO bookings "
            "(user_id,username,client_name,phone,slot_date,slot_time,service_id,comment) "
            "VALUES (?,?,?,?,?,?,?,?)",
            (uid, uname, name, phone, sd, st, service_id, comment)
        )
        bid = cur.lastrowid
        c.execute("UPDATE slots SET is_booked=1 WHERE slot_date=? AND slot_time=?", (sd, st))
    return bid

def cancel_booking_by_id(bid):
    with db() as c:
        row = c.execute(
            "SELECT slot_date,slot_time FROM bookings WHERE id=?", (bid,)
        ).fetchone()
        if row:
            c.execute(
                "UPDATE slots SET is_booked=0 WHERE slot_date=? AND slot_time=?",
                (row["slot_date"], row["slot_time"])
            )
            c.execute("DELETE FROM bookings WHERE id=?", (bid,))
    cancel_reminders(bid)

def cancel_user_booking_by_id(bid, uid):
    """Скасовує конкретний запис користувача."""
    with db() as c:
        row = c.execute(
            "SELECT * FROM bookings WHERE id=? AND user_id=?", (bid, uid)
        ).fetchone()
        if row:
            c.execute(
                "UPDATE slots SET is_booked=0 WHERE slot_date=? AND slot_time=?",
                (row["slot_date"], row["slot_time"])
            )
            c.execute("DELETE FROM bookings WHERE id=?", (bid,))
            cancel_reminders(bid)
            return dict(row)
    return None

def add_slot(sd, st, allowed_services="*"):
    try:
        with db() as c:
            c.execute(
                "INSERT INTO slots (slot_date,slot_time,allowed_services) VALUES (?,?,?) "
                "ON CONFLICT(slot_date,slot_time) DO UPDATE SET allowed_services=excluded.allowed_services",
                (sd, st, allowed_services)
            )
        return True
    except Exception as e:
        logger.error(e); return False

def get_slot_services(sd: str, st: str) -> list:
    """Повертає список дозволених services для слоту. [] = всі."""
    with db() as c:
        row = c.execute(
            "SELECT allowed_services FROM slots WHERE slot_date=? AND slot_time=?", (sd, st)
        ).fetchone()
    if not row or row["allowed_services"] == "*":
        return []  # порожній список = всі послуги дозволені
    return [s.strip() for s in row["allowed_services"].split(",") if s.strip()]

def is_service_allowed(sd: str, st: str, service_id: str) -> bool:
    allowed = get_slot_services(sd, st)
    return not allowed or service_id in allowed  # якщо порожній = всі дозволені

def remove_slot_db(sd, st):
    with db() as c:
        row = c.execute(
            "SELECT is_booked FROM slots WHERE slot_date=? AND slot_time=?", (sd, st)
        ).fetchone()
        if not row or row["is_booked"]: return False
        c.execute("DELETE FROM slots WHERE slot_date=? AND slot_time=?", (sd, st))
    return True

def close_day_db(sd):
    with db() as c:
        c.execute("UPDATE slots SET is_closed=1 WHERE slot_date=? AND is_booked=0", (sd,))

def open_day_db(sd):
    with db() as c:
        c.execute("UPDATE slots SET is_closed=0 WHERE slot_date=?", (sd,))

# ── Налаштування нагадувань ────────────────────────────────────
def get_reminder_key(uid: int) -> str:
    with db() as c:
        row = c.execute(
            "SELECT reminder_key FROM user_settings WHERE user_id=?", (uid,)
        ).fetchone()
    return row["reminder_key"] if row else "1h"

def set_reminder_key(uid: int, key: str):
    with db() as c:
        c.execute(
            "INSERT INTO user_settings (user_id, reminder_key) VALUES (?,?) "
            "ON CONFLICT(user_id) DO UPDATE SET reminder_key=excluded.reminder_key",
            (uid, key)
        )

# ================================================================
# ПЛАНУВАЛЬНИК
# ================================================================
scheduler = AsyncIOScheduler(timezone="Europe/Kyiv")

async def send_reminder_before(bot: Bot, uid: int, bid: int, st: str, mins: int):
    """Нагадування ДО запису."""
    try:
        time_str = f"{mins // 60} год" if mins >= 60 else f"{mins} хв"
        await bot.send_message(
            uid,
            f"⏰ <b>Нагадування — через {time_str} ваш запис!</b>\n\n"
            f"Сьогодні о <b>{st}</b> у вас манікюр/педікюр.\n\nДо зустрічі! 💅",
            parse_mode=ParseMode.HTML
        )
        with db() as c:
            c.execute("UPDATE bookings SET reminder_before_sent=1 WHERE id=?", (bid,))
        logger.info(f"✅ rem_before sent: bid={bid}")
    except Exception as e:
        logger.error(f"rem_before: {e}")

async def send_reminder_after(bot: Bot, uid: int, bid: int, name: str):
    """Нагадування ПІСЛЯ — через 30 днів."""
    try:
        await bot.send_message(
            uid,
            f"💅 <b>Час оновити манікюр!</b>\n\n"
            f"Привіт, <b>{name}</b>! Минув місяць від вашого останнього візиту.\n\n"
            f"Натисніть кнопку щоб записатися знову 👇",
            parse_mode=ParseMode.HTML,
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📅 Записатися", callback_data="book_start")]
            ])
        )
        with db() as c:
            c.execute("UPDATE bookings SET reminder_after_sent=1 WHERE id=?", (bid,))
        logger.info(f"✅ rem_after sent: bid={bid}")
    except Exception as e:
        logger.error(f"rem_after: {e}")

def schedule_reminders(bot: Bot, bid: int, uid: int, sd: str, st: str,
                        client_name: str, reminder_key: str):
    now = datetime.now()
    try:
        visit_dt = datetime.strptime(f"{sd} {st}", "%Y-%m-%d %H:%M")
        mins     = REMINDER_MINUTES.get(reminder_key, 60)

        if mins > 0:
            t = visit_dt - timedelta(minutes=mins)
            if t > now:
                scheduler.add_job(
                    send_reminder_before,
                    trigger=DateTrigger(run_date=t),
                    args=[bot, uid, bid, st, mins],
                    id=f"rem_before_{bid}", replace_existing=True
                )
                logger.info(f"⏰ rem_before scheduled: {t} bid={bid}")
            else:
                logger.info(f"⏭ rem_before skipped (past): bid={bid}")

        t = visit_dt + timedelta(days=30)
        if t > now:
            scheduler.add_job(
                send_reminder_after,
                trigger=DateTrigger(run_date=t),
                args=[bot, uid, bid, client_name],
                id=f"rem_after_{bid}", replace_existing=True
            )
            logger.info(f"📅 rem_after scheduled: {t} bid={bid}")

    except Exception as e:
        logger.error(f"schedule_reminders: {e}")

def cancel_reminders(bid: int):
    for jid in (f"rem_before_{bid}", f"rem_after_{bid}"):
        try: scheduler.remove_job(jid)
        except Exception: pass

def restore_reminders(bot: Bot):
    with db() as c:
        s_cols = [r[1] for r in c.execute("PRAGMA table_info(slots)").fetchall()]
        if "allowed_services" not in s_cols:
            c.execute("ALTER TABLE slots ADD COLUMN allowed_services TEXT DEFAULT '*'")
            logger.info("Міграція slots: allowed_services")

        us_cols = [r[1] for r in c.execute("PRAGMA table_info(user_settings)").fetchall()]
        if "reminder_key" in us_cols:
            rows = c.execute(
                "SELECT b.id, b.user_id, b.slot_date, b.slot_time, b.client_name, "
                "b.reminder_before_sent, b.reminder_after_sent, "
                "COALESCE(us.reminder_key,'1h') as rem_key "
                "FROM bookings b LEFT JOIN user_settings us ON us.user_id=b.user_id"
            ).fetchall()
        else:
            rows = c.execute(
                "SELECT id, user_id, slot_date, slot_time, client_name, "
                "reminder_before_sent, reminder_after_sent, '1h' as rem_key "
                "FROM bookings"
            ).fetchall()

    now = datetime.now()
    n   = 0
    for r in rows:
        try:
            vdt  = datetime.strptime(f"{r['slot_date']} {r['slot_time']}", "%Y-%m-%d %H:%M")
            mins = REMINDER_MINUTES.get(r["rem_key"], 60)

            if not r["reminder_before_sent"] and mins > 0:
                t = vdt - timedelta(minutes=mins)
                if t > now:
                    scheduler.add_job(
                        send_reminder_before, trigger=DateTrigger(run_date=t),
                        args=[bot, r["user_id"], r["id"], r["slot_time"], mins],
                        id=f"rem_before_{r['id']}", replace_existing=True)
                    n += 1

            if not r["reminder_after_sent"]:
                t = vdt + timedelta(days=30)
                if t > now:
                    scheduler.add_job(
                        send_reminder_after, trigger=DateTrigger(run_date=t),
                        args=[bot, r["user_id"], r["id"], r["client_name"]],
                        id=f"rem_after_{r['id']}", replace_existing=True)
                    n += 1
        except Exception as e:
            logger.error(f"restore bid={r['id']}: {e}")

    logger.info(f"🔄 Відновлено нагадувань: {n}")

# ================================================================
# ПІДПИСКА
# ================================================================
async def is_subscribed(bot: Bot, uid: int) -> bool:
    if not CHANNEL_ID:
        return True
    try:
        m = await bot.get_chat_member(CHANNEL_ID, uid)
        return m.status not in ("left", "kicked", "banned")
    except Exception as e:
        logger.warning(f"sub: {e}"); return True

# ================================================================
# ДОПОМІЖНЕ
# ================================================================
MONTHS_UA  = ["","Січень","Лютий","Березень","Квітень","Травень","Червень",
               "Липень","Серпень","Вересень","Жовтень","Листопад","Грудень"]
MONTHS_GEN = ["","Січня","Лютого","Березня","Квітня","Травня","Червня",
               "Липня","Серпня","Вересня","Жовтня","Листопада","Грудня"]
DAYS_UA    = ["Пн","Вт","Ср","Чт","Пт","Сб","Нд"]
DAYS_FULL  = ["Понеділок","Вівторок","Середа","Четвер","П'ятниця","Субота","Неділя"]

ALL_TIMES = [
    "08:00","08:30","09:00","09:30","10:00","10:30",
    "11:00","11:30","12:00","12:30","13:00","13:30",
    "14:00","14:30","15:00","15:30","16:00","16:30",
    "17:00","17:30","18:00","18:30","19:00","19:30","20:00"
]

def fmt_date(d_str: str) -> str:
    d = datetime.strptime(d_str, "%Y-%m-%d")
    return f"{DAYS_FULL[d.weekday()]}, {d.day} {MONTHS_GEN[d.month]}"

def is_admin(uid: int) -> bool:
    return uid == ADMIN_ID

def service_label(sid: str) -> str:
    s = SERVICE_MAP.get(sid)
    return f"{s[1]} — {s[2]} ₴" if s else (sid or "—")

async def ds(target, text: str, reply_markup=None) -> Message:
    """Видаляє старе повідомлення і надсилає нове."""
    if isinstance(target, CallbackQuery):
        bot = target.bot; chat_id = target.message.chat.id
        try: await target.message.delete()
        except TelegramBadRequest: pass
    else:
        bot = target.bot; chat_id = target.chat.id
        try: await target.delete()
        except TelegramBadRequest: pass
    return await bot.send_message(
        chat_id, text, reply_markup=reply_markup, parse_mode=ParseMode.HTML
    )

# ================================================================
# КЛАВІАТУРИ
# ================================================================
def kb_main():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📅 Записатися",            callback_data="book_start")],
        [InlineKeyboardButton(text="📋 Мої записи",            callback_data="my_bookings")],
        [InlineKeyboardButton(text="🔔 Нагадування",           callback_data="rem_settings")],
        [InlineKeyboardButton(text="💰 Прайси",                callback_data="prices")],
        [InlineKeyboardButton(text="🖼 Портфоліо",             callback_data="portfolio")],
    ])

def kb_back():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")]
    ])

def kb_sub():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📢 Підписатися на канал", url=CHANNEL_LINK)],
        [InlineKeyboardButton(text="✅ Перевірити підписку",  callback_data="check_sub")],
    ])

def kb_admin():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Додати робочий день",     callback_data="adm_add")],
        [InlineKeyboardButton(text="🗑 Видалити слот",           callback_data="adm_rm")],
        [InlineKeyboardButton(text="🔒 Закрити / відкрити день", callback_data="adm_close")],
        [InlineKeyboardButton(text="❌ Скасувати запис клієнта", callback_data="adm_cancel")],
        [InlineKeyboardButton(text="📋 Переглянути розклад",     callback_data="adm_view")],
        [InlineKeyboardButton(text="📊 Вивантажити записи",      callback_data="adm_export")],
        [InlineKeyboardButton(text="⬅️ Головне меню",            callback_data="back_to_menu")],
    ])

def kb_booking_mode():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💅 Спочатку обрати послугу",    callback_data="book_mode_a")],
        [InlineKeyboardButton(text="📅 Спочатку обрати дату і час", callback_data="book_mode_b")],
        [InlineKeyboardButton(text="⬅️ Назад до меню",              callback_data="back_to_menu")],
    ])

def kb_confirm_another(existing_info: str):
    """Підтвердження повторного запису коли вже є активний."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Так, записатися ще раз", callback_data="book_anyway")],
        [InlineKeyboardButton(text="📋 Мої записи",             callback_data="my_bookings")],
        [InlineKeyboardButton(text="⬅️ Назад до меню",          callback_data="back_to_menu")],
    ])

def kb_services(back_cb: str):
    rows = []
    for sid, name, price in SERVICES:
        rows.append([InlineKeyboardButton(
            text=f"{name} — {price} ₴",
            callback_data=f"svc|{sid}"
        )])
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=back_cb)])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_services_filtered(back_cb: str, allowed: list):
    """Показує тільки дозволені послуги. allowed=[] означає всі."""
    rows = []
    for sid, name, price in SERVICES:
        if allowed and sid not in allowed:
            continue  # пропускаємо недозволені
        rows.append([InlineKeyboardButton(
            text=f"{name} — {price} ₴",
            callback_data=f"svc|{sid}"
        )])
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=back_cb)])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_reminder_settings(current_key: str):
    rows = []
    for key, label in REMINDER_OPTIONS:
        mark = "✅ " if key == current_key else ""
        rows.append([InlineKeyboardButton(
            text=f"{mark}{label}",
            callback_data=f"set_rem|{key}"
        )])
    rows.append([InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_skip_comment(back_cb: str = "back_to_menu"):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➡️ Пропустити",   callback_data="skip_comment")],
        [InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")],
    ])

def kb_confirm():
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Підтвердити", callback_data="confirm_booking"),
        InlineKeyboardButton(text="❌ Скасувати",  callback_data="back_to_menu"),
    ]])

def kb_calendar(year: int, month: int, highlighted: list,
                cb_prefix: str, back_cb: str):
    today = date.today()
    total = monthrange(year, month)[1]
    start = date(year, month, 1).weekday()
    rows  = []

    rows.append([InlineKeyboardButton(
        text=f"📅  {MONTHS_UA[month]}  {year}", callback_data="noop"
    )])
    rows.append([InlineKeyboardButton(text=d, callback_data="noop") for d in DAYS_UA])

    week = [InlineKeyboardButton(text=" ", callback_data="noop")] * start
    for day in range(1, total + 1):
        d_obj = date(year, month, day)
        d_str = d_obj.strftime("%Y-%m-%d")
        if d_obj < today:
            btn = InlineKeyboardButton(text="·", callback_data="noop")
        elif d_str in highlighted:
            btn = InlineKeyboardButton(text=f"✅{day}", callback_data=f"{cb_prefix}|{d_str}")
        else:
            btn = InlineKeyboardButton(text=str(day), callback_data="noop")
        week.append(btn)
        if len(week) == 7:
            rows.append(week); week = []
    if week:
        week += [InlineKeyboardButton(text=" ", callback_data="noop")] * (7 - len(week))
        rows.append(week)

    pm = month - 1 if month > 1 else 12
    py = year      if month > 1 else year - 1
    nm = month + 1 if month < 12 else 1
    ny = year      if month < 12 else year + 1
    min_m = date(today.year, today.month, 1)
    max_m = min_m + timedelta(days=62)
    nav   = []
    if date(py, pm, 1) >= min_m:
        nav.append(InlineKeyboardButton(
            text="◀️", callback_data=f"cal_nav|{cb_prefix}|{py}|{pm}|{back_cb}"))
    else:
        nav.append(InlineKeyboardButton(text=" ", callback_data="noop"))
    nav.append(InlineKeyboardButton(text=" ", callback_data="noop"))
    if date(ny, nm, 1) <= max_m:
        nav.append(InlineKeyboardButton(
            text="▶️", callback_data=f"cal_nav|{cb_prefix}|{ny}|{nm}|{back_cb}"))
    else:
        nav.append(InlineKeyboardButton(text=" ", callback_data="noop"))
    rows.append(nav)
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=back_cb)])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_times(slots: list, slot_date: str, cb_prefix: str, back_cb: str):
    rows = []
    row  = []
    for i, t in enumerate(slots, 1):
        row.append(InlineKeyboardButton(
            text=f"🕐 {t}", callback_data=f"{cb_prefix}|{slot_date}|{t}"
        ))
        if i % 3 == 0:
            rows.append(row); row = []
    if row: rows.append(row)
    rows.append([InlineKeyboardButton(text="◀️ Назад до дат",  callback_data=back_cb)])
    rows.append([InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_admin_times(slot_date: str, selected: list):
    """Вибір часів для робочого дня."""
    rows = []
    row  = []
    for i, t in enumerate(ALL_TIMES, 1):
        mark = "✅" if t in selected else "➕"
        row.append(InlineKeyboardButton(
            text=f"{mark} {t}", callback_data=f"adm_toggle|{slot_date}|{t}"
        ))
        if i % 3 == 0:
            rows.append(row); row = []
    if row: rows.append(row)
    rows.append([InlineKeyboardButton(text="➡️ Далі: обрати послуги", callback_data="adm_times_next")])
    rows.append([InlineKeyboardButton(text="⬅️ Адмін-панель",         callback_data="admin_panel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def kb_admin_slot_services(slot_date: str, times: list,
                            current_time_idx: int, selected_svc: list):
    """
    Вибір послуг для конкретного часу.
    selected_svc = [] означає "всі послуги".
    """
    t = times[current_time_idx]
    rows = []

    # Кнопка "Всі послуги"
    all_mark = "✅ " if not selected_svc else ""
    rows.append([InlineKeyboardButton(
        text=f"{all_mark}🌟 Всі послуги",
        callback_data=f"adm_svc_all|{slot_date}|{current_time_idx}"
    )])

    # Кнопки послуг
    for sid, name, price in SERVICES:
        mark = "✅ " if sid in selected_svc else ""
        rows.append([InlineKeyboardButton(
            text=f"{mark}{name} — {price} ₴",
            callback_data=f"adm_svc_toggle|{slot_date}|{current_time_idx}|{sid}"
        )])

    # Навігація
    nav = []
    if current_time_idx < len(times) - 1:
        nav.append(InlineKeyboardButton(
            text=f"➡️ Наступний ({times[current_time_idx + 1]})",
            callback_data=f"adm_svc_next|{slot_date}|{current_time_idx}"
        ))
    else:
        nav.append(InlineKeyboardButton(
            text="💾 Зберегти все",
            callback_data=f"adm_svc_save|{slot_date}"
        ))
    rows.append(nav)

    # Прогрес
    progress = f"{current_time_idx + 1}/{len(times)}"
    rows.append([InlineKeyboardButton(
        text=f"⏰ {t}  [{progress}]",
        callback_data="noop"
    )])
    rows.append([InlineKeyboardButton(text="⬅️ Адмін-панель", callback_data="admin_panel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

# ================================================================
# РОУТЕР
# ================================================================
router = Router()

@router.message(Command("start"))
async def cmd_start(msg: Message, state: FSMContext):
    await state.clear()
    await ds(msg,
        "💅 <b>Вітаємо у боті для запису на манікюр та педікюр!</b>\n\n"
        "Оберіть потрібну дію:", kb_main()
    )

@router.callback_query(F.data == "back_to_menu")
async def cb_back_menu(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await ds(call, "💅 <b>Головне меню</b>\n\nОберіть потрібну дію:", kb_main())

@router.callback_query(F.data == "noop")
async def cb_noop(call: CallbackQuery):
    await call.answer()

@router.callback_query(F.data == "prices")
async def cb_prices(call: CallbackQuery):
    lines = ["💰 <b>Прайс-лист</b>\n"]
    for sid, name, price in SERVICES:
        lines.append(f"   {name} — <b>{price} ₴</b>")
    lines.append("\n<i>Ціни уточнюйте у майстра.</i>")
    await ds(call, "\n".join(lines), kb_back())

@router.callback_query(F.data == "portfolio")
async def cb_portfolio(call: CallbackQuery):
    await ds(call, "🖼 <b>Портфоліо майстра</b>\n\nНатисніть кнопку нижче:",
        InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎨 Переглянути портфоліо",
                                  url="https://www.pinterest.com/pin/376332112634273130/")],
            [InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")]
        ])
    )

@router.callback_query(F.data == "check_sub")
async def cb_check_sub(call: CallbackQuery, bot: Bot):
    if await is_subscribed(bot, call.from_user.id):
        await ds(call, "✅ <b>Підписку підтверджено!</b>\n\nТепер ви можете записатися.", kb_main())
    else:
        await call.answer("❌ Ви ще не підписані на канал!", show_alert=True)

# ── Нагадування ───────────────────────────────────────────────
@router.callback_query(F.data == "rem_settings")
async def cb_rem_settings(call: CallbackQuery):
    key = get_reminder_key(call.from_user.id)
    await ds(call,
        "🔔 <b>Нагадування до запису</b>\n\n"
        "Оберіть за скільки часу надсилати нагадування:\n\n"
        "<i>Нагадування через 1 місяць після візиту — завжди автоматично.</i>",
        kb_reminder_settings(key)
    )

@router.callback_query(F.data.startswith("set_rem|"))
async def cb_set_rem(call: CallbackQuery):
    key = call.data.split("|", 1)[1]
    set_reminder_key(call.from_user.id, key)
    label = dict(REMINDER_OPTIONS).get(key, key)
    await call.answer(f"✅ {label}")
    await ds(call,
        "🔔 <b>Нагадування до запису</b>\n\n"
        "Оберіть за скільки часу надсилати нагадування:\n\n"
        "<i>Нагадування через 1 місяць після візиту — завжди автоматично.</i>",
        kb_reminder_settings(key)
    )

# ── Навігація по місяцях ──────────────────────────────────────
@router.callback_query(F.data.startswith("cal_nav|"))
async def cb_cal_nav(call: CallbackQuery):
    parts     = call.data.split("|")
    cb_prefix = parts[1]
    y, m      = int(parts[2]), int(parts[3])
    back_cb   = parts[4]
    today     = date.today()

    if cb_prefix in ("bk_date_a", "bk_date_b"):
        dates = get_available_dates()
    elif cb_prefix == "adm_date_add":
        dates = [(today + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(0, 61)]
    elif cb_prefix in ("adm_date_rm", "adm_date_close", "adm_date_view"):
        dates = get_all_slot_dates()
    elif cb_prefix == "adm_date_cancel":
        dates = get_booking_dates()
    else:
        dates = []

    try:
        await call.message.edit_reply_markup(
            reply_markup=kb_calendar(y, m, dates, cb_prefix, back_cb)
        )
    except TelegramBadRequest: pass

# ================================================================
# ЗАПИС — ВИБІР РЕЖИМУ
# ================================================================
@router.callback_query(F.data == "book_start")
async def cb_book_start(call: CallbackQuery, state: FSMContext, bot: Bot):
    await call.answer()

    if not await is_subscribed(bot, call.from_user.id):
        await ds(call,
            "🔒 <b>Для запису необхідно підписатися на канал!</b>\n\n"
            "Підпишіться та натисніть «Перевірити підписку».", kb_sub())
        return

    # Перевіряємо чи є майбутні записи
    existing = get_user_active_bookings(call.from_user.id)
    if existing:
        lines = ["⚠️ <b>У вас вже є активний запис:</b>\n"]
        for b in existing:
            svc = service_label(b["service_id"])
            lines.append(f"📅 {b['slot_date']}  ⏰ {b['slot_time']}  💅 {svc}")
        lines.append("\nВи впевнені що хочете записатися ще раз?")
        await state.set_state(Booking.confirm_another)
        await ds(call, "\n".join(lines), kb_confirm_another(""))
        return

    await ds(call, "📋 <b>Як бажаєте записатися?</b>", kb_booking_mode())

@router.callback_query(F.data == "book_anyway")
async def cb_book_anyway(call: CallbackQuery, state: FSMContext):
    """Користувач підтвердив що хоче записатися повторно."""
    await ds(call, "📋 <b>Як бажаєте записатися?</b>", kb_booking_mode())

# ================================================================
# ШЛЯХ А: послуга → дата → час
# ================================================================
@router.callback_query(F.data == "book_mode_a")
async def cb_mode_a(call: CallbackQuery, state: FSMContext):
    await state.set_state(Booking.choose_service)
    await ds(call, "💅 <b>Оберіть послугу:</b>", kb_services("book_start"))

@router.callback_query(F.data.startswith("svc|"), Booking.choose_service)
async def cb_svc_a(call: CallbackQuery, state: FSMContext):
    sid = call.data.split("|", 1)[1]
    if sid not in SERVICE_MAP:
        await call.answer("Невідома послуга"); return
    await state.update_data(service_id=sid, mode="a")
    today     = date.today()
    available = get_available_dates(service_id=sid)
    if not available:
        await ds(call, "😔 <b>Вільних дат для цієї послуги немає.</b>\n\nСпробуйте пізніше.", kb_back()); return
    await state.set_state(Booking.choose_date_a)
    sname = SERVICE_MAP[sid][1]
    await ds(call,
        f"✅ Обрано: <b>{sname}</b>\n\n📅 <b>Оберіть дату:</b>",
        kb_calendar(today.year, today.month, available, "bk_date_a", "book_mode_a")
    )

@router.callback_query(F.data.startswith("bk_date_a|"))
async def cb_date_a(call: CallbackQuery, state: FSMContext):
    slot_date = call.data.split("|", 1)[1]
    slots     = get_free_slots(slot_date)
    if not slots:
        await call.answer("На цю дату немає вільних слотів!", show_alert=True); return
    await state.update_data(slot_date=slot_date)
    await state.set_state(Booking.choose_time_a)
    data  = await state.get_data()
    sid   = data["service_id"]
    sname = SERVICE_MAP[sid][1]
    # Фільтруємо слоти по послузі
    slots = get_free_slots(slot_date, sid)
    if not slots:
        await call.answer("На цю дату немає слотів для цієї послуги!", show_alert=True); return
    await ds(call,
        f"💅 <b>{sname}</b>\n📅 <b>{fmt_date(slot_date)}</b>\n\n⏰ <b>Оберіть час:</b>",
        kb_times(slots, slot_date, "bk_time_a", "bk_back_date_a")
    )

@router.callback_query(F.data == "bk_back_date_a")
async def cb_back_date_a(call: CallbackQuery, state: FSMContext):
    today     = date.today()
    data      = await state.get_data()
    sid       = data.get("service_id", "")
    available = get_available_dates(service_id=sid)
    sname     = SERVICE_MAP.get(sid, ("","",""))[1]
    await state.set_state(Booking.choose_date_a)
    await ds(call,
        f"✅ Обрано: <b>{sname}</b>\n\n📅 <b>Оберіть дату:</b>",
        kb_calendar(today.year, today.month, available, "bk_date_a", "book_mode_a")
    )

@router.callback_query(F.data.startswith("bk_time_a|"))
async def cb_time_a(call: CallbackQuery, state: FSMContext):
    _, slot_date, slot_time = call.data.split("|")
    await state.update_data(slot_date=slot_date, slot_time=slot_time)
    await state.set_state(Booking.enter_name)
    data  = await state.get_data()
    sname = SERVICE_MAP[data["service_id"]][1]
    await ds(call,
        f"✅ <b>{sname}</b>\n"
        f"📅 <b>{fmt_date(slot_date)}</b> о <b>{slot_time}</b>\n\n"
        f"✏️ Введіть ваше <b>ім'я та прізвище</b>:"
    )

# ================================================================
# ШЛЯХ Б: дата → час → послуга
# ================================================================
@router.callback_query(F.data == "book_mode_b")
async def cb_mode_b(call: CallbackQuery, state: FSMContext):
    today     = date.today()
    available = get_available_dates()
    if not available:
        await ds(call, "😔 <b>Вільних дат немає.</b>\n\nСпробуйте пізніше.", kb_back()); return
    await state.update_data(mode="b")
    await state.set_state(Booking.choose_date_b)
    await ds(call,
        "📅 <b>Оберіть дату запису:</b>\n\n✅ — є вільні слоти",
        kb_calendar(today.year, today.month, available, "bk_date_b", "book_start")
    )

@router.callback_query(F.data.startswith("bk_date_b|"))
async def cb_date_b(call: CallbackQuery, state: FSMContext):
    slot_date = call.data.split("|", 1)[1]
    slots     = get_free_slots(slot_date)
    if not slots:
        await call.answer("На цю дату немає вільних слотів!", show_alert=True); return
    await state.update_data(slot_date=slot_date)
    await state.set_state(Booking.choose_time_b)
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b>\n\n⏰ <b>Оберіть зручний час:</b>",
        kb_times(slots, slot_date, "bk_time_b", "bk_back_date_b")
    )

@router.callback_query(F.data == "bk_back_date_b")
async def cb_back_date_b(call: CallbackQuery, state: FSMContext):
    today     = date.today()
    available = get_available_dates()
    await state.set_state(Booking.choose_date_b)
    await ds(call,
        "📅 <b>Оберіть дату запису:</b>\n\n✅ — є вільні слоти",
        kb_calendar(today.year, today.month, available, "bk_date_b", "book_start")
    )

@router.callback_query(F.data.startswith("bk_time_b|"))
async def cb_time_b(call: CallbackQuery, state: FSMContext):
    _, slot_date, slot_time = call.data.split("|")
    await state.update_data(slot_date=slot_date, slot_time=slot_time)
    await state.set_state(Booking.choose_service_b)
    # Показуємо тільки послуги доступні для цього слоту
    allowed = get_slot_services(slot_date, slot_time)  # [] = всі
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b> о <b>{slot_time}</b>\n\n💅 <b>Оберіть послугу:</b>",
        kb_services_filtered(f"bk_date_b|{slot_date}", allowed)
    )

@router.callback_query(F.data.startswith("svc|"), Booking.choose_service_b)
async def cb_svc_b(call: CallbackQuery, state: FSMContext):
    sid = call.data.split("|", 1)[1]
    if sid not in SERVICE_MAP:
        await call.answer("Невідома послуга"); return
    await state.update_data(service_id=sid)
    data  = await state.get_data()
    sname = SERVICE_MAP[sid][1]
    await state.set_state(Booking.enter_name)
    await ds(call,
        f"✅ <b>{sname}</b>\n"
        f"📅 <b>{fmt_date(data['slot_date'])}</b> о <b>{data['slot_time']}</b>\n\n"
        f"✏️ Введіть ваше <b>ім'я та прізвище</b>:"
    )

# ================================================================
# СПІЛЬНЕ: ім'я → телефон → коментар → підтвердження
# ================================================================
@router.message(Booking.enter_name)
async def msg_name(msg: Message, state: FSMContext):
    name = (msg.text or "").strip()
    if len(name) < 2:
        await ds(msg, "⚠️ Введіть коректне ім'я (мінімум 2 символи):"); return
    await state.update_data(client_name=name)
    await state.set_state(Booking.enter_phone)
    await ds(msg,
        f"👤 <b>{name}</b>\n\n"
        f"📞 Введіть <b>номер телефону</b>:\n<i>Приклад: +380XXXXXXXXX</i>"
    )

@router.message(Booking.enter_phone)
async def msg_phone(msg: Message, state: FSMContext):
    phone = (msg.text or "").strip()
    if len(phone) < 7:
        await ds(msg, "⚠️ Введіть коректний номер телефону:"); return
    await state.update_data(phone=phone)
    await state.set_state(Booking.enter_comment)
    await ds(msg,
        f"📞 <b>{phone}</b>\n\n"
        f"💬 Залиште <b>коментар до запису</b> (побажання, алергії, нотатки)\n"
        f"або натисніть <b>Пропустити</b>:",
        kb_skip_comment()
    )

@router.message(Booking.enter_comment)
async def msg_comment(msg: Message, state: FSMContext):
    comment = (msg.text or "").strip()
    await _show_confirm(msg, state, comment)

@router.callback_query(F.data == "skip_comment")
async def cb_skip_comment(call: CallbackQuery, state: FSMContext):
    await _show_confirm(call, state, "")

async def _show_confirm(target, state: FSMContext, comment: str):
    """Показує підтвердження запису."""
    await state.update_data(comment=comment)
    data    = await state.get_data()
    rem_key = get_reminder_key(
        target.from_user.id if isinstance(target, (Message, CallbackQuery)) else 0
    )
    rem_lbl = dict(REMINDER_OPTIONS).get(rem_key, rem_key)
    sname   = service_label(data.get("service_id", ""))
    note    = f"🔔 Нагадування: {rem_lbl}" if rem_key != "off" else "🔕 Нагадування вимкнено"
    comment_line = f"\n💬 <b>Коментар:</b> {comment}" if comment else ""
    await state.set_state(Booking.confirm)
    await ds(target,
        f"📋 <b>Підтвердіть запис:</b>\n\n"
        f"💅 <b>{sname}</b>\n"
        f"👤 <b>{data['client_name']}</b>\n"
        f"📞 <b>{data['phone']}</b>\n"
        f"📅 <b>{data['slot_date']}</b>  ⏰ <b>{data['slot_time']}</b>"
        f"{comment_line}\n\n"
        f"{note}",
        kb_confirm()
    )

@router.callback_query(F.data == "confirm_booking")
async def cb_confirm(call: CallbackQuery, state: FSMContext, bot: Bot):
    data = await state.get_data()
    await state.clear()

    if data.get("slot_time") not in get_free_slots(data.get("slot_date", "")):
        await ds(call, "😔 <b>Цей час вже зайнятий.</b>\n\nОберіть інший.", kb_main())
        return

    rem_key = get_reminder_key(call.from_user.id)
    comment = data.get("comment", "")
    bid = create_booking(
        call.from_user.id, call.from_user.username,
        data["client_name"], data["phone"],
        data["slot_date"], data["slot_time"],
        data.get("service_id", "man_french"),
        comment
    )
    schedule_reminders(bot, bid, call.from_user.id,
                        data["slot_date"], data["slot_time"],
                        data["client_name"], rem_key)

    try: export_to_excel()
    except Exception as e: logger.error(f"Excel: {e}")

    rem_lbl      = dict(REMINDER_OPTIONS).get(rem_key, rem_key)
    note         = f"🔔 Нагадування: {rem_lbl}" if rem_key != "off" else "🔕 Нагадування вимкнено"
    sname        = service_label(data.get("service_id", ""))
    comment_line = f"\n💬 {comment}" if comment else ""

    await ds(call,
        f"🎉 <b>Запис підтверджено!</b>\n\n"
        f"💅 {sname}\n"
        f"📅 {data['slot_date']}  ⏰ {data['slot_time']}"
        f"{comment_line}\n\n"
        f"{note}\n"
        f"💅 Через місяць нагадаємо записатися знову",
        kb_main()
    )

    comment_admin = f"\n💬 {comment}" if comment else ""
    try:
        await bot.send_message(ADMIN_ID,
            f"🔔 <b>Новий запис!</b>\n\n"
            f"💅 {sname}\n"
            f"👤 {data['client_name']}\n"
            f"📞 <code>{data['phone']}</code>\n"
            f"📅 {data['slot_date']}  ⏰ {data['slot_time']}"
            f"{comment_admin}\n"
            f"🆔 <code>{call.from_user.id}</code> @{call.from_user.username or '—'}"
        )
    except Exception as e: logger.error(f"admin: {e}")

    if SCHEDULE_CHANNEL_ID:
        try:
            await bot.send_message(SCHEDULE_CHANNEL_ID,
                f"📋 <b>Новий запис</b>\n"
                f"💅 {sname}\n"
                f"📅 {data['slot_date']}  ⏰ {data['slot_time']}\n"
                f"👤 {data['client_name']}"
            )
        except Exception as e: logger.error(f"channel: {e}")

# ================================================================
# МОЇ ЗАПИСИ
# ================================================================
@router.callback_query(F.data == "my_bookings")
async def cb_my_bookings(call: CallbackQuery, state: FSMContext):
    await state.clear()
    bookings = get_user_active_bookings(call.from_user.id)
    if not bookings:
        await ds(call, "📭 <b>У вас немає активних записів.</b>", kb_back()); return

    lines = ["📋 <b>Ваші записи:</b>\n"]
    rows  = []
    for b in bookings:
        svc = service_label(b["service_id"])
        comment_line = f"  💬 {b['comment']}" if b["comment"] else ""
        lines.append(f"📅 {b['slot_date']}  ⏰ {b['slot_time']}\n💅 {svc}{comment_line}")
        rows.append([InlineKeyboardButton(
            text=f"❌ {b['slot_date']} {b['slot_time']}",
            callback_data=f"cancel_booking|{b['id']}"
        )])
    rows.append([InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")])

    await ds(call,
        "\n\n".join(lines) + "\n\nНатисніть ❌ щоб скасувати конкретний запис:",
        InlineKeyboardMarkup(inline_keyboard=rows)
    )

@router.callback_query(F.data.startswith("cancel_booking|"))
async def cb_cancel_booking(call: CallbackQuery):
    bid    = int(call.data.split("|", 1)[1])
    result = cancel_user_booking_by_id(bid, call.from_user.id)
    try: export_to_excel()
    except Exception as e: logger.error(f"Excel: {e}")
    if result:
        await call.answer("✅ Запис скасовано")
        # Оновлюємо список
        bookings = get_user_active_bookings(call.from_user.id)
        if not bookings:
            await ds(call, "✅ <b>Запис скасовано.</b>\n\nСлот знову вільний.", kb_main())
            return
        lines = ["📋 <b>Ваші записи:</b>\n"]
        rows  = []
        for b in bookings:
            svc = service_label(b["service_id"])
            comment_line = f"  💬 {b['comment']}" if b["comment"] else ""
            lines.append(f"📅 {b['slot_date']}  ⏰ {b['slot_time']}\n💅 {svc}{comment_line}")
            rows.append([InlineKeyboardButton(
                text=f"❌ {b['slot_date']} {b['slot_time']}",
                callback_data=f"cancel_booking|{b['id']}"
            )])
        rows.append([InlineKeyboardButton(text="⬅️ Назад до меню", callback_data="back_to_menu")])
        try:
            await call.message.edit_text(
                "\n\n".join(lines) + "\n\nНатисніть ❌ щоб скасувати конкретний запис:",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
                parse_mode=ParseMode.HTML
            )
        except TelegramBadRequest:
            pass
    else:
        await call.answer("❌ Запис не знайдено", show_alert=True)

# ================================================================
# АДМІН-ПАНЕЛЬ
# ================================================================
@router.message(Command("admin"))
async def cmd_admin(msg: Message, state: FSMContext):
    await state.clear()
    if not is_admin(msg.from_user.id):
        await ds(msg, "⛔ Доступ заборонено!"); return
    await ds(msg, "🔧 <b>Адмін-панель</b>\n\nОберіть дію:", kb_admin())

@router.callback_query(F.data == "admin_panel")
async def cb_admin_panel(call: CallbackQuery, state: FSMContext):
    await state.clear()
    if not is_admin(call.from_user.id):
        await call.answer("⛔ Доступ заборонено!", show_alert=True); return
    await ds(call, "🔧 <b>Адмін-панель</b>\n\nОберіть дію:", kb_admin())

@router.callback_query(F.data == "adm_add")
async def cb_adm_add(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    await state.set_state(Admin.add_pick_date)
    today      = date.today()
    all_dates  = [(today + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(0, 61)]
    await ds(call, "➕ <b>Додати робочий день</b>\n\nОберіть дату (включно з сьогодні):",
             kb_calendar(today.year, today.month, all_dates, "adm_date_add", "admin_panel"))

@router.callback_query(F.data.startswith("adm_date_add|"))
async def cb_adm_date_add(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    slot_date = call.data.split("|", 1)[1]
    slots     = get_all_slots(slot_date)
    existing  = [r["slot_time"] for r in slots]
    # Зберігаємо існуючі налаштування послуг
    existing_sv = {}
    with db() as c:
        for r in c.execute(
            "SELECT slot_time, allowed_services FROM slots WHERE slot_date=?", (slot_date,)
        ).fetchall():
            sv = r["allowed_services"] or "*"
            existing_sv[r["slot_time"]] = [] if sv == "*" else [s.strip() for s in sv.split(",")]
    await state.update_data(add_date=slot_date, add_times=existing.copy(), slot_services=existing_sv)
    await state.set_state(Admin.add_pick_times)
    note = ""
    if slot_date == date.today().strftime("%Y-%m-%d"):
        note = f"\n<i>Сьогодні, зараз: {datetime.now().strftime('%H:%M')}</i>"
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b>{note}\n\n"
        f"➕ — додати   ✅ — прибрати\n"
        f"Після вибору часів натисніть <b>➡️ Далі: обрати послуги</b>:",
        kb_admin_times(slot_date, existing)
    )

@router.callback_query(F.data.startswith("adm_toggle|"))
async def cb_adm_toggle(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    _, slot_date, slot_time = call.data.split("|")
    data  = await state.get_data()
    times = data.get("add_times", [])
    if slot_time in times: times.remove(slot_time)
    else: times.append(slot_time); times.sort()
    await state.update_data(add_times=times)
    try:
        await call.message.edit_reply_markup(reply_markup=kb_admin_times(slot_date, times))
    except TelegramBadRequest: pass
    await call.answer()

@router.callback_query(F.data == "adm_times_next")
async def cb_adm_times_next(call: CallbackQuery, state: FSMContext):
    """Перехід до вибору послуг після вибору часів."""
    if not is_admin(call.from_user.id): return
    data  = await state.get_data()
    times = data.get("add_times", [])
    slot_date = data.get("add_date", "")
    if not times:
        await call.answer("Оберіть хоча б один час!", show_alert=True); return
    # slot_services = {time: [service_ids]} або [] = всі
    slot_services = data.get("slot_services", {})
    await state.update_data(slot_services=slot_services)
    await state.set_state(Admin.add_pick_services)
    t = times[0]
    current_svc = slot_services.get(t, [])
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b>\n\n"
        f"Оберіть <b>послуги</b> для часу <b>{t}</b>:\n\n"
        f"<i>«Всі послуги» — клієнт може обрати будь-яку</i>",
        kb_admin_slot_services(slot_date, times, 0, current_svc)
    )

@router.callback_query(F.data.startswith("adm_svc_all|"))
async def cb_adm_svc_all(call: CallbackQuery, state: FSMContext):
    """Встановити 'всі послуги' для поточного часу."""
    if not is_admin(call.from_user.id): return
    parts     = call.data.split("|")
    slot_date = parts[1]
    idx       = int(parts[2])
    data      = await state.get_data()
    times     = data.get("add_times", [])
    sv        = data.get("slot_services", {})
    sv[times[idx]] = []   # порожній = всі
    await state.update_data(slot_services=sv)
    await call.answer("✅ Всі послуги")
    try:
        await call.message.edit_reply_markup(
            reply_markup=kb_admin_slot_services(slot_date, times, idx, [])
        )
    except TelegramBadRequest: pass

@router.callback_query(F.data.startswith("adm_svc_toggle|"))
async def cb_adm_svc_toggle(call: CallbackQuery, state: FSMContext):
    """Вмикає/вимикає конкретну послугу для часу."""
    if not is_admin(call.from_user.id): return
    parts     = call.data.split("|")
    slot_date = parts[1]
    idx       = int(parts[2])
    sid       = parts[3]
    data      = await state.get_data()
    times     = data.get("add_times", [])
    sv        = data.get("slot_services", {})
    t         = times[idx]
    current   = sv.get(t, [])
    if sid in current:
        current.remove(sid)
    else:
        current.append(sid)
    # Якщо вибрані всі 4 — прирівнюємо до "всіх"
    if len(current) == len(SERVICES):
        current = []
    sv[t] = current
    await state.update_data(slot_services=sv)
    await call.answer()
    try:
        await call.message.edit_reply_markup(
            reply_markup=kb_admin_slot_services(slot_date, times, idx, current)
        )
    except TelegramBadRequest: pass

@router.callback_query(F.data.startswith("adm_svc_next|"))
async def cb_adm_svc_next(call: CallbackQuery, state: FSMContext):
    """Переходить до наступного часу."""
    if not is_admin(call.from_user.id): return
    parts     = call.data.split("|")
    slot_date = parts[1]
    idx       = int(parts[2]) + 1
    data      = await state.get_data()
    times     = data.get("add_times", [])
    sv        = data.get("slot_services", {})
    t         = times[idx]
    current   = sv.get(t, [])
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b>\n\n"
        f"Оберіть <b>послуги</b> для часу <b>{t}</b>:\n\n"
        f"<i>«Всі послуги» — клієнт може обрати будь-яку</i>",
        kb_admin_slot_services(slot_date, times, idx, current)
    )

@router.callback_query(F.data.startswith("adm_svc_save|"))
async def cb_adm_svc_save(call: CallbackQuery, state: FSMContext):
    """Зберігає всі часи зі своїми послугами."""
    if not is_admin(call.from_user.id): return
    data      = await state.get_data()
    slot_date = data.get("add_date", "")
    times     = data.get("add_times", [])
    sv        = data.get("slot_services", {})
    await state.clear()

    with db() as c:
        c.execute("DELETE FROM slots WHERE slot_date=? AND is_booked=0", (slot_date,))

    added = 0
    lines = []
    for t in times:
        svc_list = sv.get(t, [])
        allowed  = ",".join(svc_list) if svc_list else "*"
        if add_slot(slot_date, t, allowed):
            added += 1
            if svc_list:
                svc_names = ", ".join(SERVICE_MAP[s][1].split("—")[1].strip() for s in svc_list if s in SERVICE_MAP)
                lines.append(f"🕐 {t} → {svc_names}")
            else:
                lines.append(f"🕐 {t} → всі послуги")

    summary = "\n".join(lines) if lines else "—"
    await ds(call,
        f"✅ <b>{fmt_date(slot_date)}</b>\n"
        f"Збережено слотів: <b>{added}</b>\n\n"
        f"{summary}",
        kb_admin()
    )

@router.callback_query(F.data == "adm_rm")
async def cb_adm_rm(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    dates = get_all_slot_dates()
    if not dates:
        await call.answer("Немає дат зі слотами", show_alert=True); return
    await state.set_state(Admin.rm_pick_date)
    today = date.today()
    await ds(call, "🗑 <b>Видалити слот</b>\n\nОберіть дату:",
             kb_calendar(today.year, today.month, dates, "adm_date_rm", "admin_panel"))

@router.callback_query(F.data.startswith("adm_date_rm|"))
async def cb_adm_date_rm(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    slot_date = call.data.split("|", 1)[1]
    all_slots = get_all_slots(slot_date)
    if not all_slots:
        await call.answer("Немає слотів на цю дату", show_alert=True); return

    bookings  = {b["slot_time"]: b for b in get_bookings_for_date(slot_date)}
    free      = get_free_slots(slot_date)

    # Будуємо текст з інфо по кожному слоту
    lines = [f"📅 <b>{fmt_date(slot_date)}</b>\n\n<b>Всі слоти:</b>"]
    for s in all_slots:
        t = s["slot_time"]
        if s["is_closed"]:
            lines.append(f"🔒 {t} — закрито")
        elif s["is_booked"] and t in bookings:
            b       = bookings[t]
            uname   = b["username"]
            tg_link = f'<a href="tg://user?id={b["user_id"]}">@{uname}</a>' if uname else f'<a href="tg://user?id={b["user_id"]}">#{b["user_id"]}</a>'
            lines.append(f"🔴 {t} — <b>{b['client_name']}</b> {tg_link} (зайнято, не можна видалити)")
        else:
            lines.append(f"🟢 {t} — вільно")

    if not free:
        await ds(call,
            "\n".join(lines) + "\n\n<i>Немає вільних слотів для видалення.</i>",
            kb_admin()
        ); return

    await state.update_data(rm_date=slot_date)
    await state.set_state(Admin.rm_pick_time)
    await ds(call,
        "\n".join(lines) + "\n\n🗑 <b>Оберіть вільний слот для видалення:</b>",
        kb_times(free, slot_date, "adm_rm_time", "adm_rm")
    )

@router.callback_query(F.data.startswith("adm_rm_time|"))
async def cb_adm_rm_time(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    _, slot_date, slot_time = call.data.split("|")
    ok = remove_slot_db(slot_date, slot_time)
    await state.clear()
    text = f"✅ Слот <b>{slot_date} {slot_time}</b> видалено!" if ok else "❌ Не вдалося видалити."
    await ds(call, text, kb_admin())

@router.callback_query(F.data == "adm_close")
async def cb_adm_close(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    dates = get_all_slot_dates()
    if not dates:
        await call.answer("Немає дат зі слотами", show_alert=True); return
    await state.set_state(Admin.close_pick_date)
    today = date.today()
    await ds(call, "🔒 <b>Закрити / відкрити день</b>\n\nОберіть дату:",
             kb_calendar(today.year, today.month, dates, "adm_date_close", "admin_panel"))

@router.callback_query(F.data.startswith("adm_date_close|"))
async def cb_adm_date_close(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    slot_date = call.data.split("|", 1)[1]
    await state.clear()
    with db() as c:
        row = c.execute(
            "SELECT COUNT(*) as cnt FROM slots WHERE slot_date=? AND is_closed=1", (slot_date,)
        ).fetchone()
    if row["cnt"] > 0:
        open_day_db(slot_date); text = f"🔓 День <b>{fmt_date(slot_date)}</b> відкрито!"
    else:
        close_day_db(slot_date); text = f"🔒 День <b>{fmt_date(slot_date)}</b> закрито!"
    await ds(call, text, kb_admin())

@router.callback_query(F.data == "adm_cancel")
async def cb_adm_cancel(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    dates = get_booking_dates()
    if not dates:
        await call.answer("Немає активних записів", show_alert=True); return
    await state.set_state(Admin.cancel_pick_date)
    today = date.today()
    await ds(call, "❌ <b>Скасувати запис клієнта</b>\n\nОберіть дату:",
             kb_calendar(today.year, today.month, dates, "adm_date_cancel", "admin_panel"))

@router.callback_query(F.data.startswith("adm_date_cancel|"))
async def cb_adm_date_cancel(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    slot_date = call.data.split("|", 1)[1]
    bookings  = get_bookings_for_date(slot_date)
    if not bookings:
        await call.answer("Немає записів на цю дату", show_alert=True); return
    await state.update_data(cancel_date=slot_date)
    await state.set_state(Admin.cancel_pick_booking)
    rows = [[InlineKeyboardButton(
        text=f"🕐{b['slot_time']}  👤{b['client_name']}",
        callback_data=f"adm_cancel_bid|{b['id']}"
    )] for b in bookings]
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="adm_cancel")])
    await ds(call,
        f"📅 <b>{fmt_date(slot_date)}</b>\n\nОберіть запис для скасування:",
        InlineKeyboardMarkup(inline_keyboard=rows)
    )

@router.callback_query(F.data.startswith("adm_cancel_bid|"))
async def cb_adm_cancel_bid(call: CallbackQuery, state: FSMContext, bot: Bot):
    if not is_admin(call.from_user.id): return
    bid = int(call.data.split("|", 1)[1])
    with db() as c:
        row = c.execute("SELECT * FROM bookings WHERE id=?", (bid,)).fetchone()
    if not row:
        await call.answer("Запис не знайдено", show_alert=True); return
    uid_c = row["user_id"]
    sd, st, name = row["slot_date"], row["slot_time"], row["client_name"]
    cancel_booking_by_id(bid)
    await state.clear()
    try: export_to_excel()
    except Exception as e: logger.error(f"Excel: {e}")
    try:
        await bot.send_message(uid_c,
            f"❌ <b>Ваш запис скасовано адміністратором.</b>\n\n📅 {sd}  ⏰ {st}")
    except Exception: pass
    await ds(call, f"✅ Запис <b>{name}</b> ({sd} {st}) скасовано!", kb_admin())

@router.callback_query(F.data == "adm_view")
async def cb_adm_view(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    dates = get_all_slot_dates()
    if not dates:
        await call.answer("Немає дат зі слотами", show_alert=True); return
    await state.set_state(Admin.view_pick_date)
    today = date.today()
    await ds(call, "📋 <b>Переглянути розклад</b>\n\nОберіть дату:",
             kb_calendar(today.year, today.month, dates, "adm_date_view", "admin_panel"))

@router.callback_query(F.data.startswith("adm_date_view|"))
async def cb_adm_date_view(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id): return
    slot_date = call.data.split("|", 1)[1]
    await state.clear()
    slots    = get_all_slots(slot_date)
    bookings = {b["slot_time"]: b for b in get_bookings_for_date(slot_date)}
    if not slots:
        await ds(call, f"📭 На <b>{slot_date}</b> слотів немає.", kb_admin()); return
    lines = [f"📋 <b>Розклад на {fmt_date(slot_date)}:</b>\n"]
    for s in slots:
        t = s["slot_time"]
        if s["is_closed"]:
            lines.append(f"🔒 {t} — закрито")
        elif s["is_booked"] and t in bookings:
            b       = bookings[t]
            sname   = SERVICE_MAP.get(dict(b).get("service_id",""), ("","",""))[1]
            cmnt    = f"  💬{b['comment']}" if b["comment"] else ""
            uname   = b["username"]
            tg_link = f' <a href="tg://user?id={b["user_id"]}">@{uname}</a>' if uname else f' <a href="tg://user?id={b["user_id"]}">#{b["user_id"]}</a>'
            lines.append(f"✅ {t} — <b>{b['client_name']}</b>{tg_link} 📞{b['phone']} {sname}{cmnt} [#{b['id']}]")
        else:
            lines.append(f"🟢 {t} — вільно")
    await ds(call, "\n".join(lines), kb_admin())

# ================================================================
# EXCEL ЕКСПОРТ
# ================================================================
def export_to_excel() -> str:
    with db() as c:
        rows = c.execute(
            "SELECT id, client_name, phone, slot_date, slot_time, service_id, comment, "
            "username, created_at, reminder_before_sent, reminder_after_sent "
            "FROM bookings ORDER BY slot_date, slot_time"
        ).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Записи"

    hfill  = PatternFill("solid", fgColor="C8A8E9")
    wfill  = PatternFill("solid", fgColor="FFFFFF")
    afill  = PatternFill("solid", fgColor="FAF4FF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin   = Side(style="thin", color="C8A8E9")
    brd    = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:K1")
    c1 = ws["A1"]
    c1.value = "💅 Записи на манікюр / педікюр"
    c1.font  = Font(bold=True, size=14, color="3B0066")
    c1.fill  = PatternFill("solid", fgColor="E8D5FF")
    c1.alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    c2 = ws["A2"]
    c2.value = f"Оновлено: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    c2.font  = Font(italic=True, size=9, color="888888")
    c2.alignment = center

    headers    = ["#", "Ім'я", "Телефон", "Дата", "Час", "Послуга", "Коментар",
                  "Telegram", "Запис зроблено", "⏰ Нагадування", "💅 Через місяць"]
    col_widths = [4,    22,     16,         12,     7,     28,        25,
                  18,    18,              14,              14]
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(bold=True, size=10, color="3B0066")
        cell.fill = hfill; cell.alignment = center; cell.border = brd
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[3].height = 20

    for i, row in enumerate(rows, 1):
        r    = 3 + i
        fill = afill if i % 2 == 0 else wfill
        sname = SERVICE_MAP.get(row["service_id"], ("","—",""))[1]
        vals = [
            i, row["client_name"], row["phone"], row["slot_date"], row["slot_time"],
            sname, row["comment"] or "",
            f"@{row['username']}" if row["username"] else "—", row["created_at"],
            "✅ Надіслано" if row["reminder_before_sent"] else "⏳ Очікує",
            "✅ Надіслано" if row["reminder_after_sent"]  else "⏳ Очікує",
        ]
        aligns = [center, left, left, center, center, left, left, left, left, center, center]
        for col, (val, aln) in enumerate(zip(vals, aligns), 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(size=10); cell.fill = fill
            cell.alignment = aln; cell.border = brd
        ws.row_dimensions[r].height = 18

    if not rows:
        ws.merge_cells("A4:K4")
        e = ws["A4"]; e.value = "Записів поки немає"
        e.font = Font(italic=True, color="999999"); e.alignment = center

    tr = 4 + len(rows)
    ws.merge_cells(f"A{tr}:C{tr}")
    sc = ws[f"A{tr}"]
    sc.value = f"Всього записів: {len(rows)}"
    sc.font  = Font(bold=True, size=10, color="3B0066")
    sc.fill  = PatternFill("solid", fgColor="E8D5FF")
    sc.alignment = left
    ws.freeze_panes = "A4"
    wb.save(EXCEL_FILE)
    return EXCEL_FILE

@router.message(Command("export"))
async def cmd_export(msg: Message):
    if not is_admin(msg.from_user.id):
        await ds(msg, "⛔ Доступ заборонено!"); return
    wait = await ds(msg, "⏳ Генерую файл...")
    try:
        path = export_to_excel()
        await wait.delete()
        await msg.bot.send_document(
            msg.chat.id, FSInputFile(path, filename="bookings.xlsx"),
            caption=f"📊 <b>Всі записи</b>  {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            parse_mode=ParseMode.HTML
        )
    except Exception as e: logger.error(f"export: {e}")

@router.callback_query(F.data == "adm_export")
async def cb_adm_export(call: CallbackQuery, bot: Bot):
    if not is_admin(call.from_user.id): return
    await call.answer("Генерую...")
    try:
        path = export_to_excel()
        await bot.send_document(
            call.from_user.id, FSInputFile(path, filename="bookings.xlsx"),
            caption=f"📊 <b>Всі записи</b>  {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            parse_mode=ParseMode.HTML
        )
    except Exception as e: logger.error(f"export: {e}")

# ================================================================
# ЗАПУСК
# ================================================================
async def main():
    init_db()
    bot = Bot(token=BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp  = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)
    restore_reminders(bot)
    scheduler.start()
    logger.info("🚀 Бот запущено!")
    try:
        await dp.start_polling(bot, allowed_updates=["message", "callback_query"])
    finally:
        scheduler.shutdown()
        await bot.session.close()

if __name__ == "__main__":
    asyncio.run(main())